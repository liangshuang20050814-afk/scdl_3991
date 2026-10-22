"""Create slide-ready BC1 patch-budget figures and a PowerPoint deck."""

from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.colors import TwoSlopeNorm
import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


INPUT_DIR = Path("output/comparison/ST/andersson/patch_budget")
OUTPUT_DIR = INPUT_DIR / "presentation"
METHODS = ["Random", "Similarity", "Diversity", "ST Hard"]
BUDGETS = [5, 10, 15, 20]
COLORS = {
    "Random": "#49A6FF",
    "Similarity": "#33D6B4",
    "Diversity": "#B58CFF",
    "ST Hard": "#FF8A65",
}
BG = "#07131F"
PANEL = "#102532"
TEXT = "#F4F7FA"
MUTED = "#A8BAC8"
GRID = "#36505F"
POSITIVE = "#54D6A1"
NEGATIVE = "#FF7D78"

METRICS = [
    ("CCC", "CCC", True, "Agreement in trend, mean and variance"),
    ("PCC(M)", "PCC(M)", True, "Mean correlation across 250 genes"),
    ("PCC(H)", "PCC(H)", True, "Mean correlation across top-50 genes"),
    ("MSE", "MSE", False, "Large errors count more"),
    ("MAE", "MAE", False, "Average absolute prediction error"),
]


def load_data():
    budget = pd.read_csv(INPUT_DIR / "budget_summary.csv")
    pcch = pd.read_csv(INPUT_DIR / "pcch_summary.csv")
    renamed = budget.rename(columns={
        "test_ConcordanceCorrCoef_mean": "CCC",
        "test_PearsonCorrCoef_mean": "PCC(M)",
        "test_MeanSquaredError_mean": "MSE",
        "test_MeanAbsoluteError_mean": "MAE",
    })
    pcch = pcch.rename(columns={"mean": "PCC(H)"})
    return renamed.merge(
        pcch[["method", "budget", "PCC(H)"]],
        on=["method", "budget"],
        how="left",
        validate="one_to_one",
    )


def improvement(value, baseline, higher_is_better):
    if higher_is_better:
        return (value - baseline) / abs(baseline) * 100
    return (baseline - value) / abs(baseline) * 100


def improvement_table(data):
    baseline = data[data["method"] == "Baseline"].iloc[0]
    rows = []
    for method in METHODS:
        for budget in BUDGETS:
            source = data[
                (data["method"] == method) & (data["budget"] == budget)
            ].iloc[0]
            row = {"method": method, "budget": budget}
            for column, _, higher, _ in METRICS:
                row[column] = improvement(source[column], baseline[column], higher)
                row[f"{column}_raw"] = source[column]
            rows.append(row)
    return pd.DataFrame(rows), baseline


def style_axis(axis):
    axis.set_facecolor(PANEL)
    axis.tick_params(colors=MUTED, labelsize=9)
    for spine in axis.spines.values():
        spine.set_color(GRID)
    axis.grid(axis="y", color=GRID, alpha=0.45, linewidth=0.8)


def save(fig, stem):
    png = OUTPUT_DIR / f"{stem}.png"
    pdf = OUTPUT_DIR / f"{stem}.pdf"
    # Preserve the exact 16:9 canvas so PowerPoint can place the image without
    # stretching or cropping.
    fig.savefig(png, dpi=240, facecolor=fig.get_facecolor())
    fig.savefig(pdf, facecolor=fig.get_facecolor())
    plt.close(fig)
    return png, pdf


def plot_budget_trends(improvements):
    white_bg = "#FFFFFF"
    light_panel = "#F6F8FB"
    dark_text = "#17212B"
    light_grid = "#D8E0E8"
    fig = plt.figure(figsize=(16, 9), facecolor=white_bg)
    grid = fig.add_gridspec(
        2, 3, left=0.055, right=0.97, top=0.82, bottom=0.09,
        wspace=0.23, hspace=0.42,
    )
    positions = [(0, 0), (0, 1), (0, 2), (1, 0), (1, 1)]
    all_values = improvements[[item[0] for item in METRICS]].to_numpy()
    bound = max(5.0, np.ceil(np.max(np.abs(all_values)) + 1))

    for (column, label, higher, explanation), position in zip(METRICS, positions):
        axis = fig.add_subplot(grid[position])
        axis.set_facecolor(light_panel)
        axis.tick_params(colors="#526273", labelsize=9)
        for spine in axis.spines.values():
            spine.set_color("#B8C5D0")
        axis.grid(axis="y", color=light_grid, alpha=0.85, linewidth=0.8)
        axis.axhspan(0, bound, color=POSITIVE, alpha=0.055)
        axis.axhspan(-bound, 0, color=NEGATIVE, alpha=0.045)
        axis.axhline(0, color="#344454", linestyle="--", linewidth=1.2)
        for method in METHODS:
            selected = improvements[improvements["method"] == method].sort_values(
                "budget"
            )
            axis.plot(
                selected["budget"], selected[column], marker="o", markersize=6,
                linewidth=2.25, color=COLORS[method], label=method,
            )
        axis.scatter([25], [0], marker="*", s=145, color="#F5C451", zorder=5)
        axis.set_xlim(3.5, 26.3)
        axis.set_ylim(-bound, bound)
        axis.set_xticks([5, 10, 15, 20, 25])
        axis.set_xticklabels(["5", "10", "15", "20", "25\nBaseline"])
        arrow = "↑" if higher else "↓"
        axis.set_title(
            f"{label} {arrow}\n{explanation}",
            color=dark_text,
            fontsize=12.5,
            pad=9,
        )
        axis.set_ylabel(
            "Improvement vs Baseline-25 (%)", color="#526273", fontsize=9
        )

    guide = fig.add_subplot(grid[1, 2])
    guide.set_facecolor(white_bg)
    guide.axis("off")
    guide.text(
        0.02, 0.96, "HOW TO READ", color="#68D6FF", fontsize=12,
        fontweight="bold", transform=guide.transAxes,
    )
    guide.text(
        0.02, 0.66,
        "0% = original TRIPLEX with 25 patches\nAbove 0 = better than Baseline-25\nBelow 0 = worse than Baseline-25",
        color=dark_text, fontsize=12, linespacing=1.55, transform=guide.transAxes,
    )
    handles = [
        plt.Line2D([0], [0], color=COLORS[method], marker="o", lw=2.5)
        for method in METHODS
    ]
    guide.legend(
        handles, METHODS, loc="lower left", frameon=False, ncol=2,
        labelcolor=dark_text, fontsize=10, bbox_to_anchor=(0.0, 0.14),
    )
    guide.text(
        0.02, 0.02,
        "K includes the central target patch.\nMeans across 8 patient-level folds.",
        color="#526273", fontsize=10.5, transform=guide.transAxes,
    )
    fig.text(
        0.055, 0.94, "PATCH BUDGET — PERFORMANCE ACROSS K",
        color=dark_text, fontsize=25, fontweight="bold",
    )
    fig.text(
        0.055, 0.875,
        "Arrows show the original metric direction; plotted improvement is harmonized so upward always means better",
        color="#526273", fontsize=13,
    )
    return save(fig, "01_patch_budget_trends")


def plot_heatmap(improvements):
    metric_columns = [item[0] for item in METRICS]
    values = improvements[metric_columns].to_numpy()
    limit = np.ceil(np.max(np.abs(values)))
    norm = TwoSlopeNorm(vmin=-limit, vcenter=0, vmax=limit)
    fig, axis = plt.subplots(figsize=(16, 9), facecolor=BG)
    axis.set_facecolor(BG)
    image = axis.imshow(values, cmap="RdYlGn", norm=norm, aspect="auto")
    labels = [
        f"{row.method}   K={int(row.budget)}"
        for row in improvements.itertuples(index=False)
    ]
    axis.set_xticks(range(5), [item[1] for item in METRICS], fontsize=12)
    axis.set_yticks(range(len(labels)), labels, fontsize=10)
    axis.tick_params(colors=TEXT, length=0)
    axis.xaxis.tick_top()
    for row in range(values.shape[0]):
        for column in range(values.shape[1]):
            value = values[row, column]
            axis.text(
                column, row, f"{value:+.1f}%", ha="center", va="center",
                fontsize=10.5, fontweight="bold",
                color="white" if abs(value) > limit * 0.48 else "#13232E",
            )
    for boundary in (3.5, 7.5, 11.5):
        axis.axhline(boundary, color=BG, linewidth=4)
    for spine in axis.spines.values():
        spine.set_visible(False)
    colorbar = fig.colorbar(image, ax=axis, shrink=0.77, pad=0.025)
    colorbar.ax.tick_params(colors=TEXT)
    colorbar.outline.set_edgecolor(GRID)
    colorbar.set_label("Improvement vs Baseline-25 (%)", color=TEXT, fontsize=11)
    fig.text(
        0.055, 0.955, "PATCH BUDGET — COMPLETE RESULT MAP",
        color=TEXT, fontsize=25, fontweight="bold",
    )
    fig.text(
        0.055, 0.905,
        "Green = better than Baseline-25  |  Red = worse  |  Every cell is based on the 8-fold mean",
        color=MUTED, fontsize=13,
    )
    fig.subplots_adjust(left=0.20, right=0.91, top=0.84, bottom=0.07)
    return save(fig, "02_patch_budget_heatmap")


def plot_k5_scorecard(improvements):
    subset = improvements[improvements["budget"] == 5]
    fig, axes = plt.subplots(1, 5, figsize=(16, 9), facecolor=BG, sharey=True)
    y = np.arange(len(METHODS))
    all_values = subset[[item[0] for item in METRICS]].to_numpy()
    bound = max(5.0, np.ceil(np.max(np.abs(all_values)) + 1))
    for index, (axis, (column, label, _, explanation)) in enumerate(
        zip(axes, METRICS)
    ):
        style_axis(axis)
        ordered = subset.set_index("method").loc[METHODS]
        vals = ordered[column].to_numpy()
        bar_colors = [POSITIVE if value >= 0 else NEGATIVE for value in vals]
        axis.barh(y, vals, color=bar_colors, alpha=0.92, height=0.56)
        axis.axvline(0, color="#E6EDF3", linewidth=1.2)
        axis.set_xlim(-bound, bound)
        short_explanation = {
            "CCC": "Agreement",
            "PCC(M)": "All 250 genes",
            "PCC(H)": "Top-50 genes",
            "MSE": "Squared error",
            "MAE": "Absolute error",
        }[label]
        axis.set_title(
            f"{label}\n{short_explanation}", color=TEXT, fontsize=11.5, pad=12
        )
        axis.set_xlabel("Improvement (%)", color=MUTED, fontsize=9)
        axis.grid(axis="x", color=GRID, alpha=0.45)
        axis.grid(axis="y", visible=False)
        for row, value in enumerate(vals):
            alignment = "left" if value >= 0 else "right"
            offset = 0.20 if value >= 0 else -0.20
            axis.text(
                value + offset, row, f"{value:+.1f}%", va="center", ha=alignment,
                color=TEXT, fontsize=10, fontweight="bold",
            )
        if index == 0:
            axis.set_yticks(y, METHODS, color=TEXT, fontsize=11)
        else:
            axis.tick_params(labelleft=False)
        axis.invert_yaxis()
    fig.text(
        0.055, 0.95, "K=5 SCORECARD — TARGET + FOUR CONTEXT PATCHES",
        color=TEXT, fontsize=24, fontweight="bold",
    )
    fig.text(
        0.055, 0.895,
        "A compact context is often enough, but the best selection strategy depends on the metric",
        color=MUTED, fontsize=13,
    )
    fig.text(
        0.055, 0.055,
        "At K=5: ST Hard gives the highest PCC(M); Similarity gives the highest PCC(H); Random gives the best CCC, MSE and MAE.",
        color="#DCE7EF", fontsize=12.5,
    )
    fig.subplots_adjust(left=0.10, right=0.98, top=0.80, bottom=0.14, wspace=0.24)
    return save(fig, "03_k5_scorecard")


def plot_best_settings(data, improvements, baseline):
    cards = []
    for column, label, higher, explanation in METRICS:
        candidates = data[data["method"] != "Baseline"]
        best_index = candidates[column].idxmax() if higher else candidates[column].idxmin()
        best = candidates.loc[best_index]
        gain = improvement(best[column], baseline[column], higher)
        cards.append((label, explanation, best["method"], int(best["budget"]), best[column], gain))

    fig, axis = plt.subplots(figsize=(16, 9), facecolor=BG)
    axis.axis("off")
    fig.text(0.055, 0.94, "BEST OBSERVED SETTING FOR EACH METRIC", color=TEXT,
             fontsize=25, fontweight="bold")
    fig.text(0.055, 0.885, "No single method–budget combination wins every metric",
             color=MUTED, fontsize=13)
    card_width = 0.17
    gap = 0.018
    start = 0.055
    for index, (label, explanation, method, budget, raw, gain) in enumerate(cards):
        left = start + index * (card_width + gap)
        rectangle = plt.Rectangle(
            (left, 0.22), card_width, 0.54, transform=fig.transFigure,
            facecolor=PANEL, edgecolor="#2E4B5C", linewidth=1.5,
        )
        fig.patches.append(rectangle)
        fig.text(left + 0.018, 0.695, label, color="#68D6FF", fontsize=17,
                 fontweight="bold")
        fig.text(left + 0.018, 0.625, method, color=TEXT, fontsize=16,
                 fontweight="bold")
        fig.text(left + 0.018, 0.575, f"K = {budget}", color=MUTED, fontsize=12)
        fig.text(left + 0.018, 0.455, f"{raw:.3f}", color=TEXT, fontsize=27,
                 fontweight="bold")
        fig.text(left + 0.018, 0.385, f"{gain:+.1f}% vs baseline", color=POSITIVE,
                 fontsize=12, fontweight="bold")
        fig.text(left + 0.018, 0.275, explanation, color=MUTED, fontsize=8.8,
                 wrap=True)
    fig.text(
        0.055, 0.10,
        "Interpretation: K=5 gives the strongest PCC(M) and PCC(H); Similarity K=15 gives the lowest errors.",
        color="#DCE7EF", fontsize=13,
    )
    return save(fig, "04_best_settings")


def add_textbox(slide, left, top, width, height, text, size, color, bold=False):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.name = "Arial"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = RGBColor(*color)
    return box


def build_powerpoint(images):
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]

    slide = presentation.slides.add_slide(blank)
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = RGBColor(7, 19, 31)
    add_textbox(slide, 0.72, 1.15, 11.8, 1.2, "BC1 Patch-Budget Experiment", 34, (244, 247, 250), True)
    add_textbox(slide, 0.75, 2.18, 11.1, 0.8, "How many context patches are needed—and which selection strategy works best?", 20, (168, 186, 200))
    add_textbox(slide, 0.75, 3.45, 3.2, 0.55, "5 metrics", 23, (104, 214, 255), True)
    add_textbox(slide, 4.15, 3.45, 3.2, 0.55, "4 budgets", 23, (84, 214, 161), True)
    add_textbox(slide, 7.55, 3.45, 3.2, 0.55, "4 strategies", 23, (181, 140, 255), True)
    add_textbox(slide, 0.75, 4.02, 3.2, 1.0, "CCC · PCC(M) · PCC(H)\nMSE · MAE", 13, (220, 231, 239))
    add_textbox(slide, 4.15, 4.02, 3.2, 1.0, "K = 5 · 10 · 15 · 20\nBaseline = 25", 13, (220, 231, 239))
    add_textbox(slide, 7.55, 4.02, 4.2, 1.0, "Random · Similarity\nDiversity · ST Hard", 13, (220, 231, 239))
    add_textbox(slide, 0.75, 6.55, 11.5, 0.35, "BC1 · 8-fold patient-level cross-validation · K includes the target patch", 11, (126, 151, 167))

    for image_path in images:
        slide = presentation.slides.add_slide(blank)
        slide.shapes.add_picture(
            str(image_path), 0, 0,
            width=presentation.slide_width,
            height=presentation.slide_height,
        )
    output = OUTPUT_DIR / "BC1_patch_budget_visuals.pptx"
    presentation.save(output)
    return output


def write_notes(data, improvements, baseline):
    best_lines = []
    for column, label, higher, _ in METRICS:
        candidates = data[data["method"] != "Baseline"]
        index = candidates[column].idxmax() if higher else candidates[column].idxmin()
        row = candidates.loc[index]
        gain = improvement(row[column], baseline[column], higher)
        best_lines.append(
            f"- {label}: {row['method']} K={int(row['budget'])}, "
            f"{row[column]:.3f} ({gain:+.1f}% vs Baseline-25)."
        )
    text = "\n".join([
        "# BC1 patch-budget presentation visuals",
        "",
        "All figures include CCC, PCC(M), PCC(H), MSE and MAE. Explained variance is excluded.",
        "",
        "## Suggested presentation message",
        "",
        "Reducing context from 25 patches to K=5 does not necessarily reduce prediction quality. The best correlation results occur at K=5, while Similarity K=15 gives the lowest errors. Performance is not monotonic with increasing K, suggesting that some context patches are redundant or distracting.",
        "",
        "Important wording: K=5 is sufficient among the tested budgets; it has not been proven to be the minimum possible budget.",
        "",
        "## Best observed setting",
        "",
        *best_lines,
        "",
        "## Reading the normalized charts",
        "",
        "The plotted percentage is improvement relative to Baseline-25. For CCC/PCC, higher raw values are better. For MSE/MAE, reductions are converted to positive improvements, so upward/green always means better.",
    ])
    path = OUTPUT_DIR / "speaker_notes.md"
    path.write_text(text + "\n", encoding="utf-8")
    return path


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    data = load_data()
    improvements, baseline = improvement_table(data)
    improvements.to_csv(OUTPUT_DIR / "normalized_improvements.csv", index=False)
    trend_png, _ = plot_budget_trends(improvements)
    heatmap_png, _ = plot_heatmap(improvements)
    scorecard_png, _ = plot_k5_scorecard(improvements)
    best_png, _ = plot_best_settings(data, improvements, baseline)
    powerpoint = build_powerpoint([trend_png, heatmap_png, scorecard_png, best_png])
    notes = write_notes(data, improvements, baseline)
    print(powerpoint)
    print(notes)


if __name__ == "__main__":
    main()
